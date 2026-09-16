"""The docking hand-off must contain exactly the current final molecules."""

import hashlib
import json
from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from smiles2select.export import docking, excel, selection_export
from smiles2select.selection_intelligence.action_log import ActionType, SelectionAction
from smiles2select.selection_intelligence.basket import SelectionBasket
from smiles2select.selection_intelligence.constrained_selection import (
    SelectionConstraints,
    SelectionOutcome,
)
from smiles2select.selection_intelligence.recipes import SelectionRecipe
from tests import test_export_storage

result = test_export_storage.result


def test_explicit_final_ids_override_screening_and_preserve_order(result):
    ids = result.descriptors.index.tolist()
    requested = [ids[1], ids[0]]  # first molecule failed the original screening
    frame = docking.build_frame(result, record_ids=requested)
    assert frame.access_code.tolist() == ["MOL002", "MOL001"]
    assert frame.smiles.tolist() == result.descriptors.loc[requested, "canonical_smiles"].tolist()


def test_empty_final_selection_does_not_fall_back_to_screening(result, tmp_path):
    path = tmp_path / "empty.xlsx"
    with pytest.raises(ValueError, match="empty"):
        docking.export(result, path, record_ids=[])
    assert not path.exists()


@pytest.mark.parametrize("requested, message", [([99999], "unknown"), ([1, 1], "duplicate")])
def test_explicit_final_ids_cannot_be_silently_dropped(result, requested, message):
    with pytest.raises(ValueError, match=message):
        docking.build_frame(result, record_ids=requested)


def test_invalid_selected_structure_is_reported_instead_of_omitted(result):
    record_id = result.descriptors.index[0]
    descriptors = result.descriptors.assign(
        canonical_smiles=result.descriptors.canonical_smiles.mask(
            result.descriptors.index == record_id, pd.NA
        )
    )
    with pytest.raises(ValueError, match="SMILES"):
        docking.build_frame(replace(result, descriptors=descriptors), record_ids=[record_id])


def test_all_option_exports_only_valid_molecules(result):
    frame = docking.build_frame(result, docking.DockingExportOptions(selected_only=False))
    assert frame.access_code.tolist() == ["MOL001", "MOL002", "MOL003"]


@pytest.mark.parametrize("suffix", [".xlsm", ".xls", ".txt"])
def test_unsupported_docking_format_is_rejected(result, tmp_path, suffix):
    path = tmp_path / f"docking{suffix}"
    with pytest.raises(ValueError, match="format"):
        docking.export(result, path)
    assert not path.exists()


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_final_selection_round_trips_through_docking_columns(result, tmp_path, suffix):
    requested = result.descriptors.index[:2].tolist()[::-1]
    path = docking.export(result, tmp_path / f"docking{suffix}", record_ids=requested)
    frame = pd.read_csv(path) if suffix == ".csv" else pd.read_excel(path)
    assert list(frame.columns) == ["access_code", "smiles"]
    assert frame.access_code.tolist() == ["MOL002", "MOL001"]


def test_session_docking_manifest_identifies_exact_handoff_and_shortfall(result, tmp_path):
    ids = tuple(reversed(result.descriptors.index[:2].tolist()))
    artifacts = selection_export.SessionArtifacts(
        result=result,
        basket=SelectionBasket([]),
        outcome=SelectionOutcome(ids, reasons={ids[0]: ["manually rescued"]}),
        constraints=SelectionConstraints(target_count=3),
        recipe=SelectionRecipe(target_count=3),
    )
    table, manifest = selection_export.export_docking(artifacts, tmp_path / "final.xlsx")
    audit = json.loads(manifest.read_text(encoding="utf-8"))
    assert audit["selected_record_ids"] == list(ids)
    assert audit["target_count"] == 3
    assert audit["exported_count"] == 2
    assert audit["shortfall"] == 1
    assert audit["sha256"] == hashlib.sha256(table.read_bytes()).hexdigest()
    assert audit["molecules"][0]["access_code"] == "MOL002"
    assert audit["molecules"][0]["record_id"] == ids[0]
    assert audit["molecules"][0]["selection_reasons"] == ["manually rescued"]
    assert audit["recipe"]["final_selection"]["target_count"] == 3
    assert audit["warnings"]


@pytest.mark.parametrize("export_kind", ["docking", "screening", "selection"])
def test_workbooks_preserve_formula_like_identifiers_as_literal_text(result, tmp_path, export_kind):
    record_id = int(result.descriptors.index[0])
    descriptors = result.descriptors.assign(
        molecule_id=result.descriptors.molecule_id.mask(
            result.descriptors.index == record_id, "=1+1"
        )
    )
    changed = replace(result, descriptors=descriptors)
    path = tmp_path / f"{export_kind}.xlsx"
    if export_kind == "docking":
        docking.export(changed, path, record_ids=[record_id])
        sheet, column = "molecules", "access_code"
    elif export_kind == "screening":
        excel.export(changed, path)
        sheet, column = "01_SELECTED_FINAL", "ID"
    else:
        artifacts = selection_export.SessionArtifacts(
            result=changed, basket=SelectionBasket([]),
            outcome=SelectionOutcome((record_id,)), constraints=SelectionConstraints(),
            recipe=SelectionRecipe(),
        )
        selection_export.export(artifacts, path)
        sheet, column = "FINAL_SELECTED", "ID"
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    position = [cell.value for cell in worksheet[1]].index(column) + 1
    assert worksheet.cell(2, position).data_type == "s"
    assert worksheet.cell(2, position).value == "=1+1"
    assert pd.read_excel(path, sheet_name=sheet)[column].iloc[0] == "=1+1"


def test_docking_csv_keeps_formula_like_identifier_unchanged(result, tmp_path):
    record_id = int(result.descriptors.index[0])
    descriptors = result.descriptors.assign(molecule_id="=1+1")
    path = docking.export(
        replace(result, descriptors=descriptors), tmp_path / "literal.csv", record_ids=[record_id]
    )
    assert pd.read_csv(path).access_code.tolist() == ["=1+1"]


def test_docking_manifest_includes_only_applied_action_history(result, tmp_path):
    record_id = int(result.descriptors.index[0])
    basket = SelectionBasket([])
    basket.log.record(SelectionAction(
        ActionType.ADD_TO_FINAL, (record_id,), reason="rescue justified",
        previous_state={record_id: {"selection_status": "RESERVE"}},
        new_state={record_id: {"selection_status": "FINAL_SELECTED"}},
    ))
    basket.log.record(SelectionAction(ActionType.PIN, (record_id,), reason="undone pin"))
    basket.log.undo()
    artifacts = selection_export.SessionArtifacts(
        result=result, basket=basket,
        outcome=SelectionOutcome((record_id,)), constraints=SelectionConstraints(),
        recipe=SelectionRecipe(),
    )
    _, manifest = selection_export.export_docking(artifacts, tmp_path / "history.xlsx")
    history = json.loads(manifest.read_text(encoding="utf-8"))["action_history"]
    assert len(history) == 1
    assert history[0]["action_type"] == "ADD_TO_FINAL"
    assert history[0]["reason"] == "rescue justified"
    assert history[0]["previous_state"][str(record_id)]["selection_status"] == "RESERVE"
    assert history[0]["new_state"][str(record_id)]["selection_status"] == "FINAL_SELECTED"
