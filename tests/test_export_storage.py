"""Excel export and the run database."""

from __future__ import annotations

import sqlite3
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from smiles2select.export import excel
from smiles2select.io.importer import ColumnMapping, SourceFile
from smiles2select.pipeline.config import RunConfig
from smiles2select.pipeline.runner import run
from smiles2select.storage.sqlite_store import SqliteStore
from tests.conftest import REFERENCE_SMILES

pytestmark = pytest.mark.integration

PROFILES = ("lipinski", "veber", "ghose")


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    directory = tmp_path_factory.mktemp("export")
    csv = directory / "library.csv"
    pd.DataFrame(
        [
            ("MOL001", REFERENCE_SMILES["aspirin"]),
            ("MOL002", REFERENCE_SMILES["long_alkane"]),
            ("MOL003", REFERENCE_SMILES["chalcone"]),
            ("MOL004", "not_a_smiles"),
        ],
        columns=["ID", "SMILES"],
    ).to_csv(csv, index=False)

    return run(
        RunConfig(
            sources=(
                SourceFile(path=csv, mapping=ColumnMapping(smiles="SMILES", molecule_id="ID")),
            ),
            profile_ids=PROFILES,
            alert_catalogs=("brenk",),
            n_jobs=1,
            chunk_size=2,
        )
    )


def test_detailed_export_creates_one_sheet_per_broken_rule(result, tmp_path):
    path = excel.export(result, tmp_path / "detailed.xlsx", excel.ExportOptions(detailed=True))
    names = openpyxl.load_workbook(path).sheetnames
    assert "00_SUMMARY" in names
    assert "01_SELECTED_FINAL" in names
    assert "02_EXCLUDED_FINAL" in names
    assert "03_PROFILE_MATRIX" in names
    assert "CONFIG" in names
    assert any(name.startswith("FAIL_") for name in names)
    assert any(name.endswith("_HIGH") or name.endswith("_LOW") for name in names)


def test_compact_export_uses_a_single_failure_sheet(result, tmp_path):
    path = excel.export(result, tmp_path / "compact.xlsx", excel.ExportOptions(detailed=False))
    names = openpyxl.load_workbook(path).sheetnames
    assert "RULE_FAILURES" in names
    assert not any(name.endswith("_HIGH") for name in names)


def test_every_sheet_name_is_valid_for_excel(result, tmp_path):
    path = excel.export(result, tmp_path / "names.xlsx")
    for name in openpyxl.load_workbook(path).sheetnames:
        assert len(name) <= 31
        assert not set(name) & set("[]:*?/\\")


def test_status_columns_use_short_profile_labels(result):
    frame = excel.export_frame(result)
    assert "Lipinski_Status" in frame.columns
    assert "Ghose_Status" in frame.columns


def test_excluded_sheet_explains_each_exclusion(result):
    excluded = excel.excluded_sheet(result)
    assert not excluded.empty
    assert excluded["Final_Exclusion_Reasons"].notna().all()
    invalid_row = excluded[excluded["ID"] == "MOL004"].iloc[0]
    assert "invalid" in invalid_row["Final_Exclusion_Reasons"]


def test_config_sheet_records_methods_and_versions(result):
    config = excel.config_sheet(result)
    keys = set(config["key"])
    assert "rdkit_version" in keys
    assert "config_hash" in keys
    assert "logp_method" in keys
    assert any(section.startswith("perfil:") for section in config["section"])


def test_summary_sheet_carries_the_disclaimer(result):
    summary = excel.summary_sheet(result)
    assert (summary["section"] == "Aviso").any()


def test_sheet_name_sanitisation():
    assert excel.sanitize_sheet_name("A/B:C") == "A_B_C"
    assert len(excel.sanitize_sheet_name("X" * 60)) == 31


def test_store_refuses_to_overwrite_a_foreign_database(tmp_path: Path):
    """A typo in the output path must not destroy an unrelated file."""
    foreign = tmp_path / "important.sqlite"
    connection = sqlite3.connect(foreign)
    connection.execute("CREATE TABLE payroll (id INTEGER)")
    connection.close()

    with pytest.raises(ValueError, match="refusing to overwrite"):
        SqliteStore(foreign, overwrite=True)
    assert foreign.exists()


def test_store_refuses_to_overwrite_a_non_database_file(tmp_path: Path):
    text_file = tmp_path / "notes.txt"
    text_file.write_text("important notes", encoding="utf-8")
    with pytest.raises(ValueError, match="refusing to overwrite"):
        SqliteStore(text_file, overwrite=True)
    assert text_file.read_text(encoding="utf-8") == "important notes"


def test_store_replaces_its_own_database(tmp_path: Path):
    path = tmp_path / "run.sqlite"
    with SqliteStore(path) as store:
        store.write_config({"first": "run"})
    with SqliteStore(path, overwrite=True) as store:
        assert store.read_config() == {}
